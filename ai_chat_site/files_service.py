"""上传文件的存储、文字提取，以及转换为 Gemini 输入。"""

from __future__ import annotations

import io
import json
import mimetypes
import os
import uuid
from pathlib import Path

from flask import current_app

from .db import get_db

# Gemini 能原生理解的类型：直接把字节发给模型（图表、扫描件、音视频都能看懂）
NATIVE_MIME_PREFIXES = ("image/", "audio/", "video/")
NATIVE_MIMES = {"application/pdf"}
NATIVE_IMAGE_MIMES = {"image/png", "image/jpeg", "image/webp", "image/heic", "image/heif"}

# 单次请求内联数据总预算（API 上限 20MB，留出文本余量）
INLINE_BUDGET_BYTES = 18 * 1024 * 1024
MAX_EXTRACTED_CHARS = 200_000
MAX_TEXT_PER_FILE = 60_000

_TEXT_EXTS = {
    "txt", "md", "csv", "json", "log", "html", "xml", "yaml", "yml",
    "py", "js", "ts", "java", "go", "rs", "sql", "sh",
}
_EXTRA_MIMES = {
    "heic": "image/heic",
    "m4a": "audio/mp4",
    "mov": "video/quicktime",
    "md": "text/markdown",
    "yaml": "text/yaml",
    "yml": "text/yaml",
}


def upload_dir() -> Path:
    return Path(str(current_app.config.get("UPLOAD_DIR") or "/data/uploads"))


def guess_mime(filename: str, fallback: str | None = None) -> str:
    ext = os.path.splitext(filename or "")[1].lower().lstrip(".")
    if ext in _EXTRA_MIMES:
        return _EXTRA_MIMES[ext]
    return mimetypes.guess_type(filename)[0] or fallback or "application/octet-stream"


def is_native(mime: str | None) -> bool:
    mime = (mime or "").lower()
    if mime.startswith("image/"):
        return mime in NATIVE_IMAGE_MIMES
    return mime in NATIVE_MIMES or mime.startswith(("audio/", "video/"))


def is_image(mime: str | None) -> bool:
    return (mime or "").lower().startswith("image/")


def extract_text(data: bytes, filename: str, mime: str | None) -> str:
    ext = os.path.splitext(filename or "")[1].lower().lstrip(".")
    try:
        if ext in _TEXT_EXTS or (mime or "").startswith("text/"):
            return data.decode("utf-8", errors="ignore")
        if ext == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            return "\n".join(t for t in ((p.extract_text() or "") for p in reader.pages) if t.strip())
        if ext == "docx":
            from docx import Document

            doc = Document(io.BytesIO(data))
            lines = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(lines)
        if ext == "xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"## 工作表：{ws.title}")
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row):
                        out.append(",".join("" if v is None else str(v) for v in row))
                    if len(out) > 20000:
                        break
            return "\n".join(out)
        if ext == "pptx":
            from pptx import Presentation

            prs = Presentation(io.BytesIO(data))
            out = []
            for i, slide in enumerate(prs.slides, 1):
                out.append(f"## 第 {i} 页")
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                        out.append(shape.text_frame.text)
            return "\n".join(out)
    except Exception:  # noqa: BLE001
        current_app.logger.warning("extract_text failed for %s", filename, exc_info=True)
    return ""


def save_file(*, user_id: int, data: bytes, filename: str, mime: str, source: str = "upload", extracted: str = "") -> int:
    user_dir = upload_dir() / str(int(user_id))
    user_dir.mkdir(parents=True, exist_ok=True)
    ext = os.path.splitext(filename)[1].lower()
    path = user_dir / f"{uuid.uuid4().hex}{ext}"
    path.write_bytes(data)
    db = get_db()
    cur = db.execute(
        """
        INSERT INTO uploaded_files(user_id, original_name, storage_path, mime_type, size_bytes, extracted_text, source)
        VALUES(?,?,?,?,?,?,?)
        """,
        (int(user_id), filename, str(path), mime, len(data), extracted[:MAX_EXTRACTED_CHARS], source),
    )
    db.commit()
    return int(cur.lastrowid)


def load_rows(user_id: int, file_ids: list[int]) -> dict[int, dict]:
    if not file_ids:
        return {}
    placeholders = ",".join("?" * len(file_ids))
    rows = get_db().execute(
        f"""
        SELECT id, original_name, storage_path, mime_type, size_bytes, extracted_text, source
        FROM uploaded_files WHERE user_id=? AND id IN ({placeholders})
        """,
        (int(user_id), *file_ids),
    ).fetchall()
    return {int(r["id"]): dict(r) for r in rows}


def parse_ids(raw) -> list[int]:
    if isinstance(raw, str):
        try:
            raw = json.loads(raw or "[]")
        except ValueError:
            return []
    if not isinstance(raw, list):
        return []
    return list(dict.fromkeys(int(x) for x in raw if str(x).isdigit()))


def public_info(row: dict) -> dict:
    mime = str(row.get("mime_type") or "")
    return {
        "id": int(row["id"]),
        "name": str(row.get("original_name") or "附件"),
        "mime": mime,
        "size": int(row.get("size_bytes") or 0),
        "is_image": is_image(mime),
        "source": row.get("source") or "upload",
        "url": f"/api/files/{int(row['id'])}",
    }


class InlineBudget:
    """在一次请求内分配内联字节预算：先给当前轮，再从新到旧给历史。"""

    def __init__(self, total: int = INLINE_BUDGET_BYTES):
        self.left = total

    def to_parts(self, rows: list[dict]) -> list[dict]:
        parts: list[dict] = []
        for r in rows:
            name = str(r.get("original_name") or "附件")
            mime = str(r.get("mime_type") or "")
            size = int(r.get("size_bytes") or 0)
            text = str(r.get("extracted_text") or "")
            if is_native(mime) and size <= self.left:
                try:
                    data = Path(r["storage_path"]).read_bytes()
                except OSError:
                    data = None
                if data is not None:
                    self.left -= len(data)
                    # 生成图不加文字标签：模型会模仿历史里的文字而不再出图
                    if r.get("source") != "generated":
                        parts.append({"text": f"文件：{name}"})
                    parts.append({"mime": mime, "data": data})
                    continue
            if text:
                if len(text) > MAX_TEXT_PER_FILE:
                    text = text[:MAX_TEXT_PER_FILE] + "\n…（内容过长已截断）"
                parts.append({"text": f"文件：{name}\n```\n{text}\n```"})
            else:
                parts.append({"text": f"[附件 {name} 因体积过大未随本轮发送]"})
        return parts
